# -*- coding: utf-8 -*-
"""
-------------------------------------------------
Project Name: Med-MARAG
File Name: scripts/convert_requirements_csv.py
Author: SeaOcean
Create Date: 2026-03-22
Description：需求输入集格式转换脚本
-------------------------------------------------
"""
from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class RequirementRow:
    raw: str
    seq: str = ""
    category: str = ""
    requirement: str = ""


def parse_line(raw: str) -> RequirementRow:
    line = (raw or "").strip().strip("\ufeff")
    line = line.strip().strip('"').strip("'").strip()
    parts = [p.strip() for p in line.split(",", 2)]
    row = RequirementRow(raw=raw)
    if len(parts) == 1:
        row.requirement = parts[0]
        return row
    if len(parts) == 2:
        row.seq, row.requirement = parts
        return row
    row.seq, row.category, row.requirement = parts[0], parts[1], parts[2]
    return row


def normalize_requirement(text: str) -> str:
    s = (text or "").strip()
    s = s.strip("。").strip()
    if s and not s.endswith("。") and not s.endswith((".", "!", "?", "！", "？")):
        s += "。"
    return s


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def looks_like_header(values: list[str]) -> bool:
    normalized = {value.strip().lower() for value in values if value.strip()}
    header_tokens = {
        "id",
        "seq",
        "index",
        "category",
        "requirement",
        "需求",
        "需求句子",
        "分类",
        "序号",
        "编号",
    }
    return bool(normalized & header_tokens)


def parse_structured_values(values: list[str]) -> RequirementRow:
    cleaned = [value for value in values if value]
    row = RequirementRow(raw=",".join(values))
    if not cleaned:
        return row
    if len(cleaned) == 1:
        return parse_line(cleaned[0])
    if len(cleaned) == 2:
        row.category, row.requirement = cleaned[0], cleaned[1]
        return row

    row.seq = cleaned[0]
    row.category = cleaned[1]
    row.requirement = " ".join(cleaned[2:]).strip()
    return row


def load_rows_from_text_csv(path: Path) -> list[RequirementRow]:
    raw_lines = [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return [parse_line(raw) for raw in raw_lines]


def load_rows_from_standard_csv(path: Path) -> list[RequirementRow]:
    rows: list[RequirementRow] = []
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        parsed_rows = [[normalize_cell(cell) for cell in row] for row in reader]

    if parsed_rows and looks_like_header(parsed_rows[0]):
        parsed_rows = parsed_rows[1:]

    for values in parsed_rows:
        row = parse_structured_values(values)
        if row.requirement:
            rows.append(row)
    return rows


def load_rows_from_xlsx(path: Path) -> list[RequirementRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    parsed_rows = [
        [normalize_cell(cell) for cell in row]
        for row in worksheet.iter_rows(values_only=True)
        if any(normalize_cell(cell) for cell in row)
    ]

    if parsed_rows and looks_like_header(parsed_rows[0]):
        parsed_rows = parsed_rows[1:]

    rows: list[RequirementRow] = []
    for values in parsed_rows:
        row = parse_structured_values(values)
        if row.requirement:
            rows.append(row)
    workbook.close()
    return rows


def load_rows(input_path: Path) -> list[RequirementRow]:
    suffix = input_path.suffix.lower()
    if suffix == ".xlsx":
        return load_rows_from_xlsx(input_path)
    if suffix == ".csv":
        text = input_path.read_text(encoding="utf-8").splitlines()
        # WPS导出的旧格式：每行只有一个带引号的字符串，例如 "1,分类,需求"
        if text and all(line.count(",") >= 1 and line.startswith('"') and line.endswith('"') for line in text[: min(5, len(text))]):
            return load_rows_from_text_csv(input_path)
        return load_rows_from_standard_csv(input_path)
    raise ValueError(f"Unsupported input format: {input_path.suffix}. Only .csv and .xlsx are supported.")


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert requirements CSV/XLSX into a structured dataset.")
    parser.add_argument("input_file", help="Input file path (.csv or .xlsx)")
    parser.add_argument("--out-csv", required=True, help="Output CSV path")
    parser.add_argument("--out-jsonl", required=True, help="Output JSONL path")
    args = parser.parse_args()

    in_path = Path(args.input_file).expanduser().resolve()
    out_csv = Path(args.out_csv).expanduser().resolve()
    out_jsonl = Path(args.out_jsonl).expanduser().resolve()
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    out_jsonl.parent.mkdir(parents=True, exist_ok=True)

    rows = []
    for parsed in load_rows(in_path):
        parsed.requirement = normalize_requirement(parsed.requirement)
        if parsed.requirement:
            rows.append(parsed)

    with out_csv.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["id", "category", "requirement"])
        writer.writeheader()
        for idx, row in enumerate(rows, start=1):
            rid = f"REQ{idx:03d}"
            writer.writerow(
                {
                    "id": rid,
                    "category": row.category,
                    "requirement": row.requirement,
                }
            )

    with out_jsonl.open("w", encoding="utf-8") as f:
        for idx, row in enumerate(rows, start=1):
            rid = f"REQ{idx:03d}"
            f.write(
                json.dumps(
                    {
                        "id": rid,
                        "category": row.category,
                        "requirement": row.requirement,
                        "source_line": row.raw.strip(),
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )

    print(f"converted {len(rows)} rows -> {out_csv} and {out_jsonl}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
